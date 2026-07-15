import os
import shutil
import tempfile
import unittest

os.environ["MERGEKIT_CLI_SCRIPT"] = "1"


class TestResearchParsing(unittest.TestCase):
    def setUp(self):
        self.root = tempfile.mkdtemp(prefix="research_parser_")

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_docx_sections_keep_paragraph_locators(self):
        from docx import Document
        from app.model_gateway.documents import parse_document_sections

        path = os.path.join(self.root, "report.docx")
        document = Document()
        document.add_paragraph("first finding")
        document.add_paragraph("second finding")
        document.save(path)

        self.assertEqual(
            parse_document_sections(path),
            [("paragraph", 1, "first finding"), ("paragraph", 2, "second finding")],
        )

    def test_pptx_sections_keep_slide_locators(self):
        from pptx import Presentation
        from app.model_gateway.documents import parse_document_sections

        path = os.path.join(self.root, "briefing.pptx")
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.text = "slide result"
        presentation.save(path)

        self.assertEqual(parse_document_sections(path), [("slide", 1, "slide result")])

    def test_legacy_office_formats_require_the_private_parser(self):
        from app.model_gateway.documents import DocumentParseError, parse_document_sections

        for suffix in (".doc", ".ppt"):
            with self.assertRaisesRegex(DocumentParseError, "legacy_parser_required"):
                parse_document_sections(os.path.join(self.root, "legacy" + suffix))

    def test_empty_document_is_rejected_as_no_text_layer(self):
        from docx import Document
        from app.model_gateway.documents import DocumentParseError, parse_document_sections

        path = os.path.join(self.root, "empty.docx")
        Document().save(path)

        with self.assertRaisesRegex(DocumentParseError, "no_text_layer"):
            parse_document_sections(path)


if __name__ == "__main__":
    unittest.main()
